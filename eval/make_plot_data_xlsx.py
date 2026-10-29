"""
Collect the data behind plot_acc_vs_flops.py and plot_safety_vs_flops.py
(their CSVs) into one Excel workbook: a notes sheet, a wide sheet per plot
(one row per model / precision / budget) and the long-format points as plotted.

Usage (from eval/, after both plot scripts have run):
    python make_plot_data_xlsx.py --out ../context/plot_data.xlsx
"""
import argparse
import os

import pandas as pd

# Plot outputs live in the git-ignored context/ folder.
PLOTS_DIR = os.path.join(os.path.dirname(__file__), "..", "context")
BUDGET_ORDER = ["512", "1024", "2048", "4096", "No limit"]
PRECISION_ORDER = ["BF16", "GPTQ-INT8", "GPTQ-INT4", "AWQ-W4A16", "FlatQuant-W8A8KV8", "FlatQuant-W4A4KV4"]
SKILL_COLUMNS = [("aime", "AIME"), ("math", "MATH"), ("gsm8k", "GSM8K"), ("gpqa", "GPQA"), ("avg", "Average")]
SAFETY_METRICS = [("SafeChain ASR", "safechain"), ("SafeChain FRR", "safechain"), ("WGMix ASR", "wildguard"),
                  ("WGMix FRR", "wildguard"), ("StrongREJECT", "strongreject")]
SAFETY_DATASETS = [("safechain", "SafeChain"), ("wildguard", "WGMix"), ("strongreject", "StrongREJECT")]

NOTES = [
    ("Source", "eval/plot_acc_vs_flops.py and eval/plot_safety_vs_flops.py; regenerate with "
               "eval/make_plot_data_xlsx.py after rerunning them."),
    ("Mean TFLOPs", "Mean analytical FLOPs per problem / prompt (utils/flops.py: prefill + decode backbone, "
                    "lm_head per sampled token, causal attention) / 1e12, from each response's re-tokenized "
                    "prompt and its output token count. Precision-independent: quantized variants differ from "
                    "BF16 only through tokens generated."),
    ("Mean tokens", "Mean output tokens per response."),
    ("Skill", "Accuracy % counts unparsed answers as incorrect. Average = mean over AIME, MATH, GSM8K, GPQA "
              "(TFLOPs and tokens also averaged over the four). Caps: 30000 (AIME), 8192 (MATH, GPQA), "
              "4096 (GSM8K); No limit: no budget prompt, 30000 cap."),
    ("Safety", "Judged on the final answer (after </think>). ASR: % of harmful prompts with a harmful answer "
               "(WildGuard). FRR: % of benign prompts refused. StrongREJECT: mean judge score 0-1. Lower is "
               "safer. Budget runs cap at 8192 tokens, No limit at 30000. TFLOPs/tokens are per dataset "
               "(SafeChain 1000, WildGuardMix 1699, StrongREJECT 313 prompts)."),
    ("Decoding", "Sampled at temperature 0.6, top-p 0.95, except the L1-1.5B GPTQ/AWQ skill runs (greedy)."),
    ("Missing", "Blank cells: runs not finished or not run (e.g. L1-7B BF16/GPTQ/AWQ skill; L1-7B FlatQuant "
                "SafeChain/WildGuardMix)."),
]


def tidy(df):
    df = df.copy()
    df["budget"] = df["budget"].astype(str)
    df["decoding"] = df["precision"].map(lambda p: "greedy" if p.endswith("*") else "T=0.6, top-p 0.95")
    df["precision"] = df["precision"].str.rstrip("*")
    df["mean_tflops"] = df["mean_flops"] / 1e12
    return df


def order(df):
    df = df.copy()
    df["_p"] = df["precision"].map(PRECISION_ORDER.index)
    df["_b"] = df["budget"].map(BUDGET_ORDER.index)
    return df.sort_values(["size", "_p", "_b"]).drop(columns=["_p", "_b"]).reset_index(drop=True)


def wide(df, key_col, columns, value_col, value_label, datasets):
    """One row per (size, precision, budget, decoding): the value for each
    (key, label) column, then TFLOPs and tokens for each (dataset, label)."""
    rows = []
    for keys, g in df.groupby(["size", "precision", "budget", "decoding"], sort=False):
        row = dict(zip(["Model size", "Precision", "Budget", "Decoding"], keys))
        values = g.set_index(key_col)[value_col]
        for key, label in columns:
            row[f"{label} {value_label}"] = values.get(key)
        per_ds = g.groupby("dataset")[["mean_tflops", "mean_tokens"]].first()
        for ds, label in datasets:
            row[f"{label} TFLOPs"] = per_ds["mean_tflops"].get(ds)
            row[f"{label} tokens"] = per_ds["mean_tokens"].get(ds)
        rows.append(row)
    out = pd.DataFrame(rows)
    out["_p"] = out["Precision"].map(PRECISION_ORDER.index)
    out["_b"] = out["Budget"].map(BUDGET_ORDER.index)
    return out.sort_values(["Model size", "_p", "_b"]).drop(columns=["_p", "_b"]).reset_index(drop=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=str, default=os.path.join(PLOTS_DIR, "plot_data.xlsx"))
    args = parser.parse_args()

    skill = tidy(pd.read_csv(os.path.join(PLOTS_DIR, "acc_vs_flops.csv"), dtype={"budget": str}))
    safety = tidy(pd.read_csv(os.path.join(PLOTS_DIR, "safety_vs_flops.csv"), dtype={"budget": str}))

    skill_wide = wide(skill, "dataset", SKILL_COLUMNS, "acc", "Acc %", SKILL_COLUMNS)
    safety["dataset"] = safety["metric"].map(dict(SAFETY_METRICS))
    safety_wide = wide(safety, "metric", [(m, m) for m, _ in SAFETY_METRICS], "answer_value", "(answer)",
                       SAFETY_DATASETS)

    skill_long = order(skill)[["size", "precision", "budget", "decoding", "dataset", "acc", "mean_tflops",
                               "mean_flops", "mean_tokens"]]
    safety_long = order(safety)[["size", "precision", "budget", "decoding", "metric", "answer_value",
                                 "mean_tflops", "mean_flops", "mean_tokens"]]

    sheets = [("Notes", pd.DataFrame(NOTES, columns=["Item", "Description"])),
              ("Skill vs FLOPs", skill_wide), ("Safety vs FLOPs", safety_wide),
              ("Skill points (long)", skill_long), ("Safety points (long)", safety_long)]
    with pd.ExcelWriter(args.out, engine="openpyxl") as writer:
        for name, df in sheets:
            df.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = "A2" if name == "Notes" else "E2"
            for i, col in enumerate(df.columns, start=1):
                letter = ws.cell(row=1, column=i).column_letter
                width = max([len(str(col))] + [len(str(v)) for v in df[col]])
                ws.column_dimensions[letter].width = min(width + 2, 110 if name == "Notes" else 22)
                fmt = "0.00E+00" if col == "mean_flops" else "0.000" if col == "answer_value" else "0.00"
                for cell in ws[letter][1:]:
                    if isinstance(cell.value, float):
                        cell.number_format = fmt
            if name == "Notes":
                from openpyxl.styles import Alignment
                for cell in ws["B"]:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
    print(f"Wrote {args.out}: " + ", ".join(f"{n} ({len(d)} rows)" for n, d in sheets))


if __name__ == "__main__":
    main()
